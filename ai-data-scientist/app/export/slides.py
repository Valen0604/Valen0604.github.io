"""
A document as a slide deck (.pptx), drawn with python-pptx.

One answer becomes a title slide with its main points, then one slide per chart and one
per table. Long answers are split over a few slides; the deck is a summary to present,
not the report to read, so the PDF is the place for the full text.
"""
from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import markdown as md

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
PANEL = RGBColor(0xF3, 0xF4, 0xF6)
MAX_BULLETS = 7
MAX_CHARS = 220
MAX_SLIDES_PER_SECTION = 4
TABLE_ROWS = 10
TABLE_COLS = 7


def _text(slide, left, top, width, height, text, size=18, bold=False, color=INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def _title(slide, text):
    _bar(slide)
    _text(slide, Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.9), text, size=26,
          bold=True, anchor=MSO_ANCHOR.MIDDLE)


def _footer(slide, text):
    _text(slide, Inches(0.6), H - Inches(0.5), W - Inches(1.2), Inches(0.35), text, size=9,
          color=MUTED)


def _bullets(slide, items: list[list[dict]], top=Inches(1.4)):
    box = slide.shapes.add_textbox(Inches(0.7), top, W - Inches(1.4), H - top - Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for runs in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        lead = p.add_run()
        lead.text = "•  "
        lead.font.size = Pt(16)
        lead.font.color.rgb = ACCENT
        for r in runs:
            run = p.add_run()
            run.text = r["text"]
            run.font.size = Pt(16)
            run.font.bold = bool(r.get("bold"))
            run.font.italic = bool(r.get("italic"))
            run.font.color.rgb = INK
            if r.get("code"):
                run.font.name = "Consolas"


def _clip(runs: list[dict], limit: int = MAX_CHARS) -> list[dict]:
    out, used = [], 0
    for r in runs:
        if used >= limit:
            break
        text = r["text"][:limit - used]
        used += len(text)
        out.append({**r, "text": text})
    if used >= limit:
        out.append(md.run("…"))
    return out


def _points(blocks: list[dict]) -> list[list[dict]]:
    """The sentences worth a bullet: headings, paragraphs, list items, quotes."""
    points: list[list[dict]] = []
    for b in blocks:
        if b["type"] in ("paragraph", "quote"):
            text = md.plain(b["runs"]).strip()
            if text:
                points.append(_clip(b["runs"]))
        elif b["type"] == "heading":
            points.append([{**r, "bold": True} for r in _clip(b["runs"])])
        elif b["type"] == "list":
            for it in b["items"]:
                points.append(_clip(it["runs"]))
    return points


def _picture(slide, path: Path, top=Inches(1.3)):
    try:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        w, h = 1600, 1000
    max_w, max_h = W - Inches(1.2), H - top - Inches(0.8)
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    slide.shapes.add_picture(str(path), int((W - pw) / 2), top, width=Emu(pw), height=Emu(ph))


def _table(slide, block: dict, top=Inches(1.4)):
    head = block["head"][:TABLE_COLS]
    rows = block["rows"][:TABLE_ROWS]
    ncols, nrows = max(1, len(head)), len(rows) + 1
    width = W - Inches(1.2)
    height = Inches(0.4) * nrows
    shape = slide.shapes.add_table(nrows, ncols, Inches(0.6), top, width, height)
    table = shape.table
    for c, cell in enumerate(head):
        _cell(table.cell(0, c), md.plain(cell), bold=True)
    for r, row in enumerate(rows, start=1):
        for c in range(ncols):
            _cell(table.cell(r, c), md.plain(row[c]) if c < len(row) else "")
    if len(block["rows"]) > TABLE_ROWS or len(block["head"]) > TABLE_COLS:
        _text(slide, Inches(0.6), top + height + Inches(0.1), width, Inches(0.4),
              f'Showing {min(len(block["rows"]), TABLE_ROWS)} of {len(block["rows"])} rows and '
              f'{ncols} of {len(block["head"])} columns.', size=10, color=MUTED)


def _cell(cell, text: str, bold: bool = False):
    cell.text = text[:60]
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11)
            r.font.bold = bold
            r.font.color.rgb = INK


def render(doc: dict) -> bytes:
    """The document as .pptx bytes."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    footer = doc.get("footer", "")

    # Title slide
    s = prs.slides.add_slide(blank)
    _bar(s)
    _text(s, Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(1.6), doc["title"], size=40,
          bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _text(s, Inches(0.8), Inches(3.9), W - Inches(1.6), Inches(0.6), doc.get("subtitle", ""),
          size=16, color=MUTED)
    if doc.get("datasets"):
        _text(s, Inches(0.8), Inches(4.6), W - Inches(1.6), Inches(1.2), "Data: " + "; ".join(
            f'{d["name"]} ({d["rows"]:,} rows, {d["columns_count"]} columns)'
            for d in doc["datasets"]), size=12, color=MUTED)
    _footer(s, footer)

    if doc.get("context"):
        s = prs.slides.add_slide(blank)
        _title(s, "About the data")
        _bullets(s, [[md.run(label + ". ", bold=True), md.run(value[:MAX_CHARS * 2])]
                     for label, value in doc["context"]])
        _footer(s, footer)

    for section in doc["sections"]:
        heading = section.get("heading") or doc["title"]
        points = _points(section["blocks"])
        made = 0
        for start in range(0, len(points), MAX_BULLETS):
            if made >= MAX_SLIDES_PER_SECTION:
                break
            s = prs.slides.add_slide(blank)
            _title(s, heading if made == 0 else heading + " (continued)")
            _bullets(s, points[start:start + MAX_BULLETS])
            _footer(s, footer)
            made += 1
        if not points and section["blocks"]:
            s = prs.slides.add_slide(blank)
            _title(s, heading)
            _footer(s, footer)
        for b in section["blocks"]:
            if b["type"] == "image":
                path = doc.get("images", {}).get(b["file"])
                if not path:
                    continue
                s = prs.slides.add_slide(blank)
                _title(s, b.get("title") or heading)
                _picture(s, path)
                _footer(s, footer)
            elif b["type"] == "table" and b["rows"]:
                s = prs.slides.add_slide(blank)
                _title(s, heading)
                _table(s, b)
                _footer(s, footer)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
